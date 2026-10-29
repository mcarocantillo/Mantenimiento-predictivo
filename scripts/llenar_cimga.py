"""Llena la guía de CIMGA 2027 (convertida a .docx) con el resumen del trabajo."""
import copy
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ENTRADA, SALIDA, RESUMEN = sys.argv[1], sys.argv[2], sys.argv[3]
F = "[FALTA]"

TITULO = "Reglas de Nelson y cartas de control I-MR, X̄-R y X̄-S para anticipar fallas de rodamientos"
OBJETIVO = ("Evaluar si las alarmas basadas en las ocho reglas de Nelson, validadas antes en otra industria, "
            "anticipan fallas de rodamientos, e identificar qué regla y qué carta de control (I-MR, X̄-R o X̄-S) "
            "relacionan mejor la alarma con la falla.")
CAMPOS = {  # (fila, columna) -> valor
    (3, 0): "María Carolina Cantillo Orozco", (3, 1): F, (3, 3): F,
    (4, 0): "maracc@uninorte.edu.co", (4, 3): "Colombia",
    (5, 0): "Universidad del Norte", (5, 2): "Estudiante de maestría",
    (6, 0): "Andrés Aldana", (6, 1): F, (6, 3): F,
    (7, 0): F, (7, 3): "Colombia",
    (8, 0): "Universidad del Norte [confirmar]", (8, 2): "Profesor [confirmar]",
}
TERCER_AUTOR = "Jenaan Tarabein"
CONTENIDO = [
    (0, "Introducción"),
    (1, "Fallas de rodamientos y alarmas por vibración"),
    (1, "Reglas de Nelson y su uso previo en otra industria"),
    (0, "Marco teórico"),
    (1, "Indicadores de vibración"),
    (2, "RMS, curtosis y valor pico"),
    (2, "Espectro de envolvente y frecuencias de defecto"),
    (1, "Cartas de control I-MR, X̄-R y X̄-S"),
    (1, "Las ocho reglas de Nelson"),
    (0, "Datos y método"),
    (1, "Conjunto IMS de la Universidad de Cincinnati"),
    (1, "Línea base por coeficiente de variación"),
    (1, "Construcción de las cartas y subgrupos"),
    (1, "Evaluación de las alarmas"),
    (2, "Sensibilidad, falsas alarmas y anticipación"),
    (2, "Prueba de permutación y comparaciones múltiples"),
    (0, "Resultados"),
    (1, "Reglas asociadas con la falla"),
    (1, "Correlación temporal entre alarmas y falla"),
    (1, "Comparación de las cartas"),
    (0, "Discusión y aplicación en planta"),
    (0, "Conclusiones y trabajo futuro"),
]
TEMA = "Mantenimiento predictivo"


def escribir(parrafo, texto, negrita=False):
    """Reemplaza el texto del párrafo por un solo run de 12 pt."""
    for r in list(parrafo.runs):
        r._r.getparent().remove(r._r)
    run = parrafo.add_run(texto)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = negrita
    # Misma fuente que el formulario: Arial (Times New Roman para texto asiático)
    fuentes = run._r.get_or_add_rPr().get_or_add_rFonts()
    for atributo, valor in [("w:ascii", "Arial"), ("w:hAnsi", "Arial"), ("w:cs", "Arial"),
                            ("w:eastAsia", "Times New Roman")]:
        fuentes.set(qn(atributo), valor)
    return run


def borrar(parrafo):
    parrafo._p.getparent().remove(parrafo._p)


d = docx.Document(ENTRADA)
t = d.tables[0]

celda = t.cell(2, 0)
escribir(celda.paragraphs[2], TITULO, negrita=True)

for (f, c), valor in CAMPOS.items():
    escribir(t.cell(f, c).paragraphs[1], valor)

# El formulario trae dos autores; el tercero se agrega copiando el bloque del segundo (filas 6 a 8)
ancla = t.rows[8]._tr
for tr in [copy.deepcopy(t.rows[i]._tr) for i in (6, 7, 8)]:
    ancla.addnext(tr)
    ancla = tr
for fila in (9, 10, 11):
    for celda in t.rows[fila].cells:
        if len(celda.paragraphs) > 1:
            escribir(celda.paragraphs[1], F)
etiqueta = t.cell(9, 0).paragraphs[0]
escribir(etiqueta, "Nombre del tercer autor:", negrita=True)
escribir(t.cell(9, 0).paragraphs[1], TERCER_AUTOR)

celda = t.cell(12, 0)
p = celda.paragraphs[1]
escribir(p, OBJETIVO)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

celda = t.cell(13, 0)
etiqueta = celda.paragraphs[0]
for r in etiqueta.runs[2:]:  # quita "(escribir en este espacio el resumen del trabajo)"
    r._r.getparent().remove(r._r)
p = celda.paragraphs[1]
escribir(p, open(RESUMEN, encoding="utf-8").read().strip())
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for sobra in celda.paragraphs[2:-1]:
    borrar(sobra)

celda = t.cell(14, 0)
modelos = {0: celda.paragraphs[1], 1: celda.paragraphs[2], 2: celda.paragraphs[3]}
ancla = modelos[0]._p
for nivel, texto in CONTENIDO:
    nuevo = copy.deepcopy(modelos[nivel]._p)
    ancla.addprevious(nuevo)
    escribir(docx.text.paragraph.Paragraph(nuevo, celda), texto)
for p in modelos.values():
    borrar(p)


def marcar(tabla):
    for fila in tabla.rows:
        if fila.cells[0].text.strip().rstrip(".") == TEMA:
            p = fila.cells[-1].paragraphs[0]
            escribir(p, "X", negrita=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return True
    return any(marcar(anidada) for c in tabla._cells for anidada in c.tables)


assert marcar(t), "no se encontró el tema"
d.save(SALIDA)
print("guardado", SALIDA)
